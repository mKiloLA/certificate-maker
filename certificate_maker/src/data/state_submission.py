"""Create state-submission workbooks from webinar exports."""
from builtins import object
import csv
import json
import os
import re
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from certificate_maker.src.data.ref import us_state_to_abbrev


REPORT_HEADERS = [
    "Total Hours",
    "Last Name",
    "First Name",
    "Email",
    "Phone",
    "State #1",
    "Bar Number",
    "Course #",
    "Polls",
    "Evaluation",
    "CofA Created",
    "CofA Sent",
    "Reported to State",
    "Paid",
    "Notes",
]

def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _email(value: object) -> str:
    return _clean(value).casefold()


def _read_csv_rows(path: str) -> list[list[str]]:
    with open(path, "r", newline="", encoding="utf-8-sig") as source:
        return [row for row in csv.reader(source)]


def _add_csv_sheet(workbook: Workbook, title: str, path: str) -> None:
    """Copy a source CSV into a worksheet without changing its layout."""
    worksheet = workbook.create_sheet(title=title)
    for row in _read_csv_rows(path):
        worksheet.append(row)


def _find_header(rows: Iterable[list[str]], *required: str) -> tuple[int, list[str]] | None:
    for index, row in enumerate(rows):
        if all(value in row for value in required):
            return index, row
    return None


def _poll_participants(poll_file: str) -> list[set[str]]:
    """Return the responding email addresses for each poll section."""
    rows = _read_csv_rows(poll_file)
    participants: list[set[str]] = []
    index = 0

    while index < len(rows):
        header = _find_header(rows[index:], "User Name", "Email Address")
        if header is None:
            break
        header_offset, columns = header
        header_index = index + header_offset
        email_index = columns.index("Email Address")
        responses: set[str] = set()

        for row in rows[header_index + 1:]:
            if not any(_clean(value) for value in row):
                break
            if len(row) > email_index and _email(row[email_index]):
                responses.add(_email(row[email_index]))

        participants.append(responses)
        index = header_index + 1

    return participants


def _poll_summary(email: str, participants: list[set[str]]) -> str:
    answered = [str(number) for number, people in enumerate(participants, 1) if email in people]
    if not answered:
        return "Did Not Respond"
    if len(answered) == len(participants):
        return f"1-{len(participants)}"
    return ", ".join(answered)


def _abbreviate_evaluation(value: object) -> str:
    text = _clean(value)
    if not text:
        return ""
    normalized = re.sub(r"\s+", " ", text).strip()
    words = normalized.split(" ")
    if normalized.casefold() == "excellent":
        return "Ex"
    if len(words) == 1:
        return words[0][0].upper()
    return "".join(word[0].upper() for word in words)


def _survey_responses(survey_file: str) -> dict[str, str]:
    """Return compact evaluation summaries keyed by respondent email."""
    rows = _read_csv_rows(survey_file)
    header_info = _find_header(rows, "User Name", "Email Address")
    if header_info is None:
        return {}

    header_index, columns = header_info
    email_index = columns.index("Email Address")
    response_indices = [
        index
        for index, column in enumerate(columns)
        if column.startswith("Please rate") or column.startswith("Would you recommend")
    ]
    responses: dict[str, str] = {}

    for row in rows[header_index + 1:]:
        if not any(_clean(value) for value in row):
            break
        if len(row) <= email_index:
            continue
        respondent = _email(row[email_index])
        if not respondent:
            continue
        answers = [
            _abbreviate_evaluation(row[index])
            for index in response_indices
            if index < len(row) and _clean(row[index])
        ]
        responses[respondent] = ", ".join(answers) if answers else "Did Not Respond"

    return responses


def _format_phone(value: object) -> str:
    digits = re.sub(r"\D", "", _clean(value))
    if len(digits) == 10:
        return f"{digits[:3]}-{digits[3:6]}-{digits[6:]}"
    return _clean(value)


def _split_name(value: object) -> tuple[str, str]:
    parts = _clean(value).split()
    if len(parts) < 2:
        return "", parts[0] if parts else ""
    return " ".join(parts[:-1]), parts[-1]


def _state_abbreviation(value: object) -> str:
    state = _clean(value)
    if len(state) == 2:
        return state.upper()
    return us_state_to_abbrev.get(state.title(), state)


def _certificate_date(value: object) -> str:
    date = _clean(value)
    return f"{date}, by CM" if date else ""


def _total_hours(attendee: dict[str, object]) -> str:
    total_hours = _clean(attendee.get("totalhours"))
    credits = _clean(attendee.get("credits"))
    if not total_hours or not credits:
        return ""
    return f"{total_hours}T/{credits}"


def _identifier(value: object) -> int | str:
    """Store digit-only identifiers numerically without losing leading zeroes."""
    identifier = _clean(value).lstrip("#")
    if identifier.isdigit() and not identifier.startswith("0"):
        return int(identifier)
    return identifier


def _json_rows(json_file: str) -> tuple[list[dict[str, object]], str, str]:
    with open(json_file, "r", encoding="utf-8") as source:
        document = json.load(source)

    attendees = document.get("attendees", [])
    if not attendees:
        return [], "", ""

    first = attendees[0]
    title = f"{_clean(first.get('clename'))} {_clean(first.get('overflow'))}".strip()
    date = _clean(first.get("certifieddate"))
    return attendees, title, date


def _format_state_submission_sheet(worksheet) -> None:
    """Apply the formatting used by the webinar worksheet report."""
    title_alignment = Alignment(horizontal="centerContinuous")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    centered = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    separator_fill = PatternFill(fill_type="solid", fgColor="595959")

    for column in range(1, len(REPORT_HEADERS) + 1):
        worksheet.cell(1, column).alignment = title_alignment
        worksheet.cell(2, column).alignment = title_alignment

    worksheet["A1"].font = Font(name="Arial", size=16, bold=True)
    worksheet["A2"].font = Font(name="Arial", size=14, bold=True)

    for cell in worksheet[4]:
        cell.font = Font(name="Arial", size=12, bold=True)
        cell.alignment = header_alignment

    centered_columns = {1, 6, 7, 8, 9, 10, 11, 12, 13, 14}
    for row in worksheet.iter_rows(min_row=5, max_row=worksheet.max_row, max_col=len(REPORT_HEADERS)):
        if not any(cell.value not in (None, "") for cell in row):
            for cell in row:
                cell.fill = separator_fill
            continue
        for column, cell in enumerate(row, start=1):
            cell.font = Font(name="Arial", size=12)
            cell.border = thin_border
            cell.alignment = centered if column in centered_columns else Alignment(vertical="center")
            if column in {7, 8} and isinstance(cell.value, str):
                cell.number_format = "@"

    for column_cells in worksheet.iter_cols(min_row=1, max_row=worksheet.max_row, max_col=len(REPORT_HEADERS)):
        width = max(len(_clean(cell.value)) for cell in column_cells) + 2
        worksheet.column_dimensions[column_cells[0].column_letter].width = width

    worksheet.freeze_panes = "A5"
    worksheet.auto_filter.ref = f"A4:O{worksheet.max_row}"


def create_state_submission_report(
    json_file: str,
    poll_file: str,
    survey_file: str,
    attendance_file: str,
) -> str:
    """Create and return the state-submission workbook path.

    The report columns come from the JSON, poll, and survey exports, while all
    three source CSVs are copied into the workbook for reference.
    """
    attendees, title, date = _json_rows(json_file)
    poll_participants = _poll_participants(poll_file)
    survey_responses = _survey_responses(survey_file)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "State Submissions"
    worksheet.append([title])
    worksheet.append([date])
    worksheet.append([])
    worksheet.append(REPORT_HEADERS)

    report_rows = []
    for attendee in attendees:
        first_name, last_name = _split_name(attendee.get("name"))
        respondent = _email(attendee.get("email"))
        report_rows.append([
            _total_hours(attendee),
            last_name,
            first_name,
            attendee.get("email", ""),
            _format_phone(attendee.get("phonenumber")),
            _state_abbreviation(attendee.get("state")),
            _identifier(attendee.get("barnumber")),
            _identifier(attendee.get("coursenumber")),
            _poll_summary(respondent, poll_participants),
            survey_responses.get(respondent, "Did Not Respond"),
            _certificate_date(attendee.get("certifieddate")),
            "",
            "",
            "",
            "",
        ])

    report_rows.sort(key=lambda row: (row[5].casefold(), row[1].casefold(), row[2].casefold()))
    previous_state = None
    for report_row in report_rows:
        current_state = report_row[5]
        if previous_state is not None and current_state != previous_state:
            worksheet.append([""] * len(REPORT_HEADERS))
        worksheet.append(report_row)
        previous_state = current_state

    _format_state_submission_sheet(worksheet)

    _add_csv_sheet(workbook, "Zoom Attend.", attendance_file)
    _add_csv_sheet(workbook, "Zoom Poll", poll_file)
    _add_csv_sheet(workbook, "Zoom Survey", survey_file)

    json_name = os.path.splitext(os.path.basename(json_file))[0]
    output_file = os.path.join(os.path.dirname(json_file), f"{json_name}.xlsx")
    workbook.save(output_file)
    return output_file
